"""Automated tests for the non-GUI logic (no camera or display required).

Run from the project folder:   python -m unittest discover -s tests -v
"""

from __future__ import annotations

import sys
import tempfile
import unittest
from datetime import datetime, timedelta
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import config  # noqa: E402
import utils  # noqa: E402
from attendance import AttendanceService, MarkStatus, RecognitionSession  # noqa: E402
from config import Settings  # noqa: E402
from database import AttendanceRecord, Database, DuplicatePersonError  # noqa: E402
from excel_manager import ExcelLockedError, ExcelManager  # noqa: E402
from face_engine import DetectedFace, FaceGallery  # noqa: E402


def random_unit(seed: int) -> np.ndarray:
    vector = np.random.default_rng(seed).normal(size=config.EMBEDDING_DIM)
    return (vector / np.linalg.norm(vector)).astype(np.float32)


class BaseCase(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        folder = Path(self._tmp.name)
        self.db = Database(folder / "test.db")
        self.db.initialize()
        self.excel = ExcelManager(folder / "attendance.xlsx")
        self.service = AttendanceService(self.db, self.excel)

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def excel_rows(self) -> list[tuple]:
        sheet = load_workbook(self.excel.path).active
        return list(sheet.iter_rows(min_row=2, values_only=True))


class DatabaseTests(BaseCase):
    def test_duplicate_person_id_rejected_case_insensitive(self) -> None:
        self.db.add_person("EMP001", "Rahul", "IT", "", [random_unit(1)])
        with self.assertRaises(DuplicatePersonError):
            self.db.add_person("emp001", "Other", "HR", "", [random_unit(2)])

    def test_unique_attendance_constraint(self) -> None:
        record = AttendanceRecord("EMP001", "Rahul", "IT", "2026-10-25", "09:05:00 AM")
        with self.db.transaction() as conn:
            self.assertTrue(self.db.insert_attendance(conn, record))
            self.assertFalse(self.db.insert_attendance(conn, record))
        self.assertEqual(len(self.db.query_attendance()), 1)

    def test_delete_person_and_reregister(self) -> None:
        self.db.add_person("EMP001", "Rahul", "IT", "", [random_unit(1)])
        self.assertTrue(self.db.delete_person("EMP001"))
        self.assertEqual(self.db.count_persons(), 0)
        self.db.add_person("EMP001", "Rahul", "IT", "", [random_unit(1)])
        self.assertEqual(self.db.count_persons(), 1)

    def test_embeddings_round_trip(self) -> None:
        samples = [random_unit(i) for i in range(5)]
        self.db.add_person("EMP001", "Rahul", "IT", "", samples)
        entry = self.db.load_gallery()[0]
        self.assertEqual(entry.embeddings.shape, (5, config.EMBEDDING_DIM))
        np.testing.assert_allclose(entry.embeddings[2], samples[2], atol=1e-6)
        self.assertEqual(self.db.list_persons()[0].sample_count, 5)


class AttendanceTests(BaseCase):
    def test_mark_once_per_day(self) -> None:
        base = datetime(2026, 10, 25, 9, 5, 0)
        first = self.service.mark_present("EMP001", "Rahul Patil", "IT", base)
        self.assertEqual(first.status, MarkStatus.MARKED)
        for minutes in (5, 85, 255):
            again = self.service.mark_present("EMP001", "Rahul Patil", "IT",
                                              base + timedelta(minutes=minutes))
            self.assertEqual(again.status, MarkStatus.ALREADY_PRESENT)
        rows = self.excel_rows()
        self.assertEqual(rows, [("EMP001", "Rahul Patil", "IT", "25-10-2026",
                                 "09:05:00 AM", "P")])
        self.assertEqual(len(self.db.query_attendance()), 1)

    def test_next_day_creates_new_record_and_keeps_old(self) -> None:
        day1 = datetime(2026, 10, 25, 9, 5, 0)
        self.service.mark_present("EMP001", "Rahul", "IT", day1)
        result = self.service.mark_present("EMP001", "Rahul", "IT", day1 + timedelta(days=1))
        self.assertEqual(result.status, MarkStatus.MARKED)
        self.assertEqual(len(self.excel_rows()), 2)

    def test_excel_locked_rolls_back_database(self) -> None:
        def locked(_record):
            raise ExcelLockedError("Please close attendance.xlsx and try again.")
        original = self.excel.append_record
        self.excel.append_record = locked          # simulate Excel holding the file
        result = self.service.mark_present("EMP001", "Rahul", "IT")
        self.assertEqual(result.status, MarkStatus.EXCEL_ERROR)
        self.assertIn("close attendance.xlsx", result.message)
        self.assertEqual(self.db.query_attendance(), [])   # nothing left behind
        self.excel.append_record = original         # Excel closed -> retry works
        retry = self.service.mark_present("EMP001", "Rahul", "IT")
        self.assertEqual(retry.status, MarkStatus.MARKED)
        self.assertEqual(len(self.excel_rows()), 1)

    def test_existing_excel_data_preserved_and_sync(self) -> None:
        self.excel.ensure_file()
        self.service.mark_present("EMP001", "Rahul", "IT", datetime(2026, 10, 24, 9, 0, 0))
        # A record that exists only in the database is added by sync, once.
        with self.db.transaction() as conn:
            self.db.insert_attendance(conn, AttendanceRecord(
                "EMP002", "Priya", "HR", "2026-10-24", "09:10:00 AM"))
        self.assertEqual(self.service.sync_excel(), 1)
        self.assertEqual(self.service.sync_excel(), 0)
        self.assertEqual(len(self.excel_rows()), 2)

    def test_export_creates_formatted_file(self) -> None:
        self.service.mark_present("EMP001", "Rahul", "IT", datetime(2026, 10, 25, 9, 5, 0))
        target = Path(self._tmp.name) / "export.xlsx"
        count = self.excel.export(self.db.query_attendance(), target)
        self.assertEqual(count, 1)
        sheet = load_workbook(target).active
        self.assertEqual([c.value for c in sheet[1]],
                         ["Employee ID", "Name", "Department", "Date", "Time", "Status"])


class GalleryAndSessionTests(BaseCase):
    def setUp(self) -> None:
        super().setUp()
        self.rahul, self.priya = random_unit(10), random_unit(20)
        self.db.add_person("EMP001", "Rahul", "IT", "", [self.rahul])
        self.db.add_person("EMP002", "Priya", "HR", "", [self.priya])
        self.gallery = FaceGallery()
        self.gallery.load(self.db.load_gallery())

    def test_match_unknown_and_uncertain(self) -> None:
        self.assertEqual(self.gallery.match(self.rahul, 0.4).employee_id, "EMP001")
        self.assertEqual(self.gallery.match(self.priya, 0.4).employee_id, "EMP002")
        self.assertEqual(self.gallery.match(random_unit(99), 0.4).status, "unknown")
        # Slightly below the threshold -> "Face not recognized", never a guess.
        blend = self.rahul * 0.36 + random_unit(5) * 0.93
        blend /= np.linalg.norm(blend)
        score = float(blend @ self.rahul)
        self.assertEqual(self.gallery.match(blend, score + 0.02).status, "uncertain")

    def test_ambiguous_face_is_not_guessed(self) -> None:
        midpoint = (self.rahul + self.priya)
        midpoint /= np.linalg.norm(midpoint)
        self.assertEqual(self.gallery.match(midpoint, 0.3).status, "uncertain")

    def test_session_confirms_then_marks_once_and_respects_cooldown(self) -> None:
        clock = [0.0]
        settings = Settings(confirm_frames=3, attendance_cooldown_seconds=60)
        rahul = self.rahul

        class FakeEngine:
            def detect(self, frame):
                return [DetectedFace((10, 10, 120, 120), 0.99, np.zeros(15, np.float32))]

            def embed(self, frame, face):
                return rahul

        session = RecognitionSession(FakeEngine(), self.gallery, self.service,
                                     settings, clock=lambda: clock[0])
        frame = np.zeros((480, 640, 3), np.uint8)
        outcomes = []
        for _ in range(10):
            clock[0] += 0.1
            overlays, results = session.process_frame(frame)
            outcomes.extend(results)
        self.assertEqual(overlays[0].label, "Rahul (EMP001)")
        self.assertEqual([r.status for r in outcomes], [MarkStatus.MARKED])
        clock[0] += 61                              # cooldown over -> DB says present
        for _ in range(4):
            clock[0] += 0.1
            _, results = session.process_frame(frame)
            outcomes.extend(results)
        self.assertEqual([r.status for r in outcomes],
                         [MarkStatus.MARKED, MarkStatus.ALREADY_PRESENT])
        self.assertEqual(len(self.db.query_attendance()), 1)


class UtilsTests(unittest.TestCase):
    def test_time_and_date_formats(self) -> None:
        moment = datetime(2026, 10, 25, 9, 15, 32)
        self.assertEqual(utils.format_date(moment), "25-10-2026")
        self.assertEqual(utils.format_time(moment), "09:15:32 AM")
        self.assertEqual(utils.format_time(datetime(2026, 1, 1, 0, 5, 0)), "12:05:00 AM")
        self.assertEqual(utils.format_time(datetime(2026, 1, 1, 13, 20, 1)), "01:20:01 PM")
        self.assertEqual(utils.display_to_iso("25-10-2026"), "2026-10-25")
        with self.assertRaises(ValueError):
            utils.display_to_iso("2026/10/25")

    def test_settings_are_clamped(self) -> None:
        settings = Settings(camera_index=-4, recognition_threshold=5, confirm_frames=0,
                            attendance_cooldown_seconds=1)
        settings.sanitize()
        self.assertEqual((settings.camera_index, settings.recognition_threshold,
                          settings.confirm_frames, settings.attendance_cooldown_seconds),
                         (0, 0.9, 1, 10))


if __name__ == "__main__":
    unittest.main()
