"""Excel attendance file handling (openpyxl).

The workbook is never overwritten: new rows are appended to the existing file,
and every save goes through a temporary file that atomically replaces the
original, so a failed save cannot corrupt earlier attendance data.
"""

from __future__ import annotations

import logging
import os
import threading
from datetime import date, datetime
from pathlib import Path
from typing import Iterable
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

import utils
from database import AttendanceRecord

logger = logging.getLogger(__name__)

SHEET_NAME = "Attendance"
HEADERS = ["Employee ID", "Name", "Department", "Date", "Time", "Status"]

_HEADER_FILL = PatternFill(fill_type="solid", start_color="1F4E78", end_color="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_CENTER = Alignment(horizontal="center", vertical="center")
_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


class ExcelError(Exception):
    """Excel operation failed. The message is safe to show to the user."""


class ExcelLockedError(ExcelError):
    """The workbook is open in another program (usually Microsoft Excel)."""


def _locked_message(path: Path) -> str:
    return f"Please close {path.name} and try again."


def _new_workbook() -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTER
        cell.border = _BORDER
    sheet.row_dimensions[1].height = 24
    sheet.freeze_panes = "A2"
    _finalize(sheet)
    return workbook


def _append_row(sheet, record: AttendanceRecord) -> None:
    sheet.append([record.employee_id, record.name, record.department,
                  record.display_date, record.attendance_time, record.status])
    for cell in sheet[sheet.max_row]:
        cell.alignment = _CENTER
        cell.border = _BORDER
        cell.number_format = "@"          # keep dates/times as plain text


def _finalize(sheet) -> None:
    """Adjust column widths and the auto-filter range."""
    for index, header in enumerate(HEADERS, start=1):
        letter = get_column_letter(index)
        longest = max(
            (len(str(cell.value)) for cell in sheet[letter] if cell.value is not None),
            default=len(header),
        )
        sheet.column_dimensions[letter].width = max(14, longest + 4)
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(sheet.max_row, 1)}"


def _date_text(value: object) -> str:
    """Normalise a date cell (text or real Excel date) to DD-MM-YYYY."""
    if isinstance(value, (datetime, date)):
        return value.strftime(utils.DATE_FORMAT)
    return str(value).strip() if value is not None else ""


def _row_key(employee_id: object, date_value: object) -> tuple[str, str]:
    return (str(employee_id).strip().lower(), _date_text(date_value))


class ExcelManager:
    """Reads and writes ``attendance.xlsx``."""

    def __init__(self, path: Path) -> None:
        self._path = Path(path)
        self._lock = threading.Lock()

    @property
    def path(self) -> Path:
        """Location of the attendance workbook."""
        return self._path

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------
    def ensure_file(self) -> None:
        """Create the workbook with formatted headers if it does not exist."""
        with self._lock:
            if self._path.exists():
                return
            self._save(_new_workbook())
            logger.info("Created %s", self._path.name)

    def append_record(self, record: AttendanceRecord) -> bool:
        """Append one attendance row.

        Returns False when the sheet already holds a row for the same person
        and date (duplicate protection), True when a row was added.
        Raises ExcelLockedError / ExcelError on failure.
        """
        with self._lock:
            self._check_writable()
            workbook = self._load_or_create()
            sheet = workbook[SHEET_NAME] if SHEET_NAME in workbook.sheetnames \
                else workbook.active
            self._validate_headers(sheet)
            new_key = _row_key(record.employee_id, record.display_date)
            if new_key in self._existing_keys(sheet):
                logger.info("Excel already contains attendance for %s", record.employee_id)
                return False
            _append_row(sheet, record)
            _finalize(sheet)
            self._save(workbook)
        return True

    def sync_missing(self, records: Iterable[AttendanceRecord]) -> int:
        """Add database records that are missing from the workbook.

        Used by "Open Excel" so the file always reflects the database.
        Returns the number of rows added.
        """
        with self._lock:
            self._check_writable()
            workbook = self._load_or_create()
            sheet = workbook[SHEET_NAME] if SHEET_NAME in workbook.sheetnames \
                else workbook.active
            self._validate_headers(sheet)
            existing = self._existing_keys(sheet)
            added = 0
            # Oldest first so the sheet stays in chronological order.
            for record in sorted(records, key=lambda r: r.attendance_date):
                if _row_key(record.employee_id, record.display_date) in existing:
                    continue
                _append_row(sheet, record)
                added += 1
            if added or not self._path.exists():
                _finalize(sheet)
                self._save(workbook)
        if added:
            logger.info("Synchronised %d missing attendance rows to Excel", added)
        return added

    def export(self, records: Iterable[AttendanceRecord], destination: Path) -> int:
        """Write records to a brand-new workbook at ``destination``."""
        destination = Path(destination)
        workbook = _new_workbook()
        sheet = workbook[SHEET_NAME]
        count = 0
        for record in records:
            _append_row(sheet, record)
            count += 1
        _finalize(sheet)
        try:
            destination.parent.mkdir(parents=True, exist_ok=True)
            workbook.save(destination)
        except PermissionError as exc:
            raise ExcelLockedError(_locked_message(destination)) from exc
        except OSError as exc:
            logger.error("Excel export failed: %s", exc)
            raise ExcelError(f"Could not save {destination.name}: {exc.strerror or exc}") from exc
        logger.info("Exported %d records to %s", count, destination.name)
        return count

    # ------------------------------------------------------------------
    # Internals
    # ------------------------------------------------------------------
    def _check_writable(self) -> None:
        """Fail early if Excel has the file open (Windows locks it)."""
        if not self._path.exists():
            return
        try:
            with open(self._path, "r+b"):
                pass
        except PermissionError as exc:
            logger.error("Excel file is locked: %s", self._path)
            raise ExcelLockedError(_locked_message(self._path)) from exc
        except OSError as exc:
            raise ExcelError(f"Cannot access {self._path.name}: {exc.strerror or exc}") from exc

    def _load_or_create(self) -> Workbook:
        if not self._path.exists():
            return _new_workbook()
        try:
            return load_workbook(self._path)
        except (BadZipFile, InvalidFileException, KeyError, ValueError) as exc:
            logger.error("Excel file unreadable: %s", exc)
            raise ExcelError(
                f"{self._path.name} looks damaged. Rename or move it, then try again "
                "(a new file will be created automatically)."
            ) from exc
        except PermissionError as exc:
            raise ExcelLockedError(_locked_message(self._path)) from exc
        except OSError as exc:
            raise ExcelError(f"Cannot open {self._path.name}: {exc.strerror or exc}") from exc

    def _validate_headers(self, sheet) -> None:
        found = [cell.value for cell in sheet[1][: len(HEADERS)]]
        if found != HEADERS:
            raise ExcelError(
                f"{self._path.name} does not have the expected columns. "
                "Rename or move it, then try again."
            )

    @staticmethod
    def _existing_keys(sheet) -> set[tuple[str, str]]:
        keys = set()
        for row in sheet.iter_rows(min_row=2, max_col=4, values_only=True):
            if row[0]:
                keys.add(_row_key(row[0], row[3]))
        return keys

    def _save(self, workbook: Workbook) -> None:
        """Save via a temp file, then atomically replace the real file."""
        temp_path = self._path.with_name(self._path.stem + ".tmp")
        try:
            workbook.save(temp_path)
            os.replace(temp_path, self._path)
        except PermissionError as exc:
            logger.error("Cannot write Excel file (locked): %s", exc)
            raise ExcelLockedError(_locked_message(self._path)) from exc
        except OSError as exc:
            logger.error("Cannot write Excel file: %s", exc)
            raise ExcelError(f"Could not save {self._path.name}: {exc.strerror or exc}") from exc
        finally:
            try:
                temp_path.unlink(missing_ok=True)
            except OSError:
                pass
